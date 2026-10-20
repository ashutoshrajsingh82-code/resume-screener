"""Server-side validation and text extraction for uploaded resume documents."""
from __future__ import annotations

import io
import re
import zipfile
from typing import Dict

class DocumentValidationError(ValueError):
    """A safe, user-facing validation or extraction problem for one document."""


ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MIME_TYPES: Dict[str, set[str]] = {
    "pdf": {"application/pdf", "application/x-pdf"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    "txt": {"text/plain"},
}
GENERIC_MIME_TYPES = {"", "application/octet-stream"}
MAX_EXPANDED_DOCX_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 2000
MAX_PDF_PAGES = 500


def validate_filename(filename: str) -> str:
    if not isinstance(filename, str) or not filename or len(filename) > 255:
        raise DocumentValidationError("The filename is missing or too long.")
    if filename in {".", ".."} or any(char in filename for char in ("/", "\\", "\x00")):
        raise DocumentValidationError("The filename is unsafe. Remove any folder path and try again.")
    if re.search(r'[<>:"|?*\x00-\x1f]', filename):
        raise DocumentValidationError("The filename contains unsupported characters. Rename the file and try again.")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentValidationError("Unsupported file type. Upload a PDF, DOCX, or TXT resume.")
    return extension


def validate_mime_type(extension: str, content_type: str | None) -> None:
    mime = (content_type or "").split(";", 1)[0].strip().lower()
    if mime not in GENERIC_MIME_TYPES and mime not in MIME_TYPES[extension]:
        raise DocumentValidationError(
            f"The uploaded content type does not match the .{extension} filename. Check the file and try again."
        )


def size_limit_label(size_bytes: int) -> str:
    if size_bytes >= 1024 * 1024:
        amount = size_bytes / (1024 * 1024)
        return f"{amount:g} MB"
    if size_bytes >= 1024:
        return f"{(size_bytes + 1023) // 1024} KB"
    return f"{size_bytes} bytes"


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            expanded_size = sum(member.file_size for member in members)
            if len(members) > MAX_DOCX_ENTRIES or expanded_size > MAX_EXPANDED_DOCX_BYTES:
                raise DocumentValidationError("This DOCX expands beyond the safe processing limit.")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentValidationError("This is not a valid DOCX document.")
        document = Document(io.BytesIO(data))
        sections = [paragraph.text for paragraph in document.paragraphs]
        sections.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(part for part in sections if part and part.strip())
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentValidationError("This DOCX file is malformed or cannot be opened.") from exc


def _extract_pdf(data: bytes) -> str:
    if b"%PDF-" not in data[:1024]:
        raise DocumentValidationError("This file does not contain a valid PDF header.")
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentValidationError("This PDF is password-protected. Remove the password and upload it again.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentValidationError(f"This PDF has more than {MAX_PDF_PAGES} pages and cannot be processed safely.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentValidationError("This PDF is malformed or cannot be opened.") from exc


def extract_resume_text(
    filename: str,
    content_type: str | None,
    data: bytes,
    max_file_bytes: int,
) -> str:
    """Validate bytes and return extracted text without writing the upload to disk."""
    extension = validate_filename(filename)
    validate_mime_type(extension, content_type)
    if len(data) > max_file_bytes:
        raise DocumentValidationError(f"This file exceeds the {size_limit_label(max_file_bytes)} upload limit.")
    if not data:
        raise DocumentValidationError("This file is empty. Select a non-empty resume.")

    if extension == "txt":
        if b"\x00" in data:
            raise DocumentValidationError("This TXT file contains binary data and cannot be read as text.")
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentValidationError("This TXT file is not valid UTF-8 text.") from exc
    elif extension == "docx":
        text = _extract_docx(data)
    else:
        text = _extract_pdf(data)

    if not text.strip():
        if extension == "pdf":
            raise DocumentValidationError("No selectable text was found. Scanned PDFs need OCR, which this prototype does not provide.")
        raise DocumentValidationError("No readable text was found in this file.")
    return text
