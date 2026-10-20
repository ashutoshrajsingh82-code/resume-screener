from io import BytesIO

import pytest
from docx import Document

from document_processing import DocumentValidationError, extract_resume_text


MAX_BYTES = 10 * 1024 * 1024


def pdf_with_text(text="Python engineer"):
    stream = f"BT /F1 14 Tf 20 50 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)


def docx_with_text(text="Python engineer"):
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "FastAPI"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_extracts_utf8_txt():
    assert extract_resume_text("resume.txt", "text/plain; charset=utf-8", b"\xef\xbb\xbfPython engineer", MAX_BYTES) == "Python engineer"


def test_extracts_docx_paragraphs_and_tables():
    text = extract_resume_text(
        "resume.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        docx_with_text(),
        MAX_BYTES,
    )
    assert "Python engineer" in text
    assert "FastAPI" in text


def test_extracts_text_from_pdf():
    assert "Python engineer" in extract_resume_text("resume.pdf", "application/pdf", pdf_with_text(), MAX_BYTES)


@pytest.mark.parametrize("filename, mime, data", [
    ("resume.exe", "application/octet-stream", b"anything"),
    ("../resume.txt", "text/plain", b"Python"),
    ("resume.pdf", "text/plain", pdf_with_text()),
])
def test_rejects_unsupported_unsafe_or_mismatched_uploads(filename, mime, data):
    with pytest.raises(DocumentValidationError):
        extract_resume_text(filename, mime, data, MAX_BYTES)


def test_rejects_oversized_file():
    with pytest.raises(DocumentValidationError, match="upload limit"):
        extract_resume_text("resume.txt", "text/plain", b"a" * 101, 100)


def test_rejects_empty_and_whitespace_files():
    with pytest.raises(DocumentValidationError, match="empty"):
        extract_resume_text("empty.txt", "text/plain", b"", MAX_BYTES)
    with pytest.raises(DocumentValidationError, match="No readable text"):
        extract_resume_text("blank.txt", "text/plain", b" \n\t", MAX_BYTES)


@pytest.mark.parametrize("filename, data", [
    ("broken.pdf", b"%PDF-1.4\nnot a PDF"),
    ("broken.docx", b"PK\x03\x04not a DOCX zip"),
])
def test_rejects_malformed_documents(filename, data):
    with pytest.raises(DocumentValidationError, match="malformed|valid DOCX"):
        extract_resume_text(filename, None, data, MAX_BYTES)


def test_scanned_pdf_is_reported_without_claiming_ocr_support():
    blank_pdf = pdf_with_text("")
    with pytest.raises(DocumentValidationError, match="Scanned PDFs need OCR"):
        extract_resume_text("scan.pdf", "application/pdf", blank_pdf, MAX_BYTES)
